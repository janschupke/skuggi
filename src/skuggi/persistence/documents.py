"""Safe, pure-Python text extraction from binary documents (pdf/docx/xlsx).

An engagement collects documents -- pulled from a target, handed over by a
client. The agent may need their text, but a document is also a classic code-
execution and resource-exhaustion vector. This module extracts text and nothing
else, under a strict set of guardrails:

- **No execution.** Pure-Python parsers only (pypdf, python-docx, openpyxl); no
  subprocess, no LibreOffice/soffice conversion, no shelling out. A PDF's
  embedded JavaScript or /Launch action is never run (pypdf only reads text);
  an Office file's macros are never run -- in fact a macro-bearing container is
  refused outright.
- **Macros refused.** docx/xlsx are ZIPs; a ``vbaProject.bin`` member (a renamed
  .docm/.xlsm) means the file is refused rather than parsed.
- **Resource limits.** A size cap before opening, a page/cell cap while reading,
  and a decompression-ratio + total-size check so a zip bomb cannot exhaust
  memory.
- **Read-only.** The file is only ever opened for reading; nothing is written,
  executed or marked executable.

Extracted text is still untrusted: callers redact it before it reaches the model
(the RAG retrieval chokepoint, or an on-demand read).
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from skuggi.common.logs import get_logger

log = get_logger(__name__)

# A document larger than this is refused before opening.
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
# Caps while reading, so a crafted file cannot exhaust memory/time.
MAX_PDF_PAGES = 1_000
MAX_CELLS = 200_000
# Zip-bomb guards for the Office container formats.
MAX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
_RATIO_FLOOR_BYTES = 64 * 1024  # ignore the ratio for tiny members

SUPPORTED_SUFFIXES = frozenset({".pdf", ".docx", ".xlsx"})
# Macro-enabled Office formats: refused by suffix before we even open them.
MACRO_SUFFIXES = frozenset({".docm", ".xlsm", ".pptm", ".dotm", ".xltm", ".xlsb"})
_MACRO_MEMBER = "vbaProject.bin"


class DocumentError(RuntimeError):
    """A document could not be parsed safely (unsupported, macro, bomb, oversize)."""


def is_supported(path: Path) -> bool:
    """Whether `path` is a document type this module can extract text from."""
    return path.suffix.lower() in SUPPORTED_SUFFIXES


def _check_size(path: Path) -> None:
    size = path.stat().st_size
    if size > MAX_DOCUMENT_BYTES:
        msg = f"document too large to parse safely ({size} bytes): {path.name}"
        raise DocumentError(msg)


def _inspect_zip(path: Path) -> None:
    """Refuse a macro-bearing or zip-bomb Office container before parsing it."""
    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as exc:
        msg = f"not a valid Office document: {path.name}"
        raise DocumentError(msg) from exc
    total = 0
    for info in infos:
        if info.filename.rsplit("/", 1)[-1] == _MACRO_MEMBER:
            msg = f"refusing a macro-bearing document ({_MACRO_MEMBER}): {path.name}"
            raise DocumentError(msg)
        total += info.file_size
        if (
            info.compress_size > 0
            and info.file_size >= _RATIO_FLOOR_BYTES
            and info.file_size / info.compress_size > MAX_COMPRESSION_RATIO
        ):
            msg = (
                f"refusing a document with a suspicious compression ratio: {path.name}"
            )
            raise DocumentError(msg)
    if total > MAX_UNCOMPRESSED_BYTES:
        msg = f"refusing a document that unpacks too large ({total} bytes): {path.name}"
        raise DocumentError(msg)


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader  # noqa: PLC0415 -- lazy: optional `documents` extra
    from pypdf.errors import PdfReadError  # noqa: PLC0415

    try:
        reader = PdfReader(path)
    except (PdfReadError, ValueError, OSError) as exc:
        msg = f"could not parse PDF: {path.name}"
        raise DocumentError(msg) from exc
    # An empty password unlocks many "encrypted" PDFs; a real one is refused
    # rather than prompted for (there is no interactive surface here).
    if reader.is_encrypted and reader.decrypt("") == 0:
        msg = f"refusing an encrypted PDF: {path.name}"
        raise DocumentError(msg)
    try:
        pages = reader.pages[:MAX_PDF_PAGES]
        return "\n\n".join(page.extract_text() or "" for page in pages).strip()
    except (PdfReadError, ValueError, OSError) as exc:
        msg = f"could not parse PDF: {path.name}"
        raise DocumentError(msg) from exc


def _extract_docx(path: Path) -> str:
    _inspect_zip(path)
    from docx import Document as DocxDocument  # noqa: PLC0415 -- lazy optional extra

    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = "\t".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts).strip()


def _extract_xlsx(path: Path) -> str:
    _inspect_zip(path)
    from openpyxl import load_workbook  # noqa: PLC0415 -- lazy optional extra

    # read_only streams rows; data_only takes cached values not formulae; macros
    # are not loaded (and a macro container was already refused by _inspect_zip).
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        cells = 0
        for sheet in wb.worksheets:
            parts.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(v) for v in row if v is not None]
                cells += len(row)
                if cells > MAX_CELLS:
                    msg = f"refusing a workbook with too many cells: {path.name}"
                    raise DocumentError(msg)
                if values:
                    parts.append("\t".join(values))
        return "\n".join(parts).strip()
    finally:
        wb.close()


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
}


def extract_text(path: Path) -> str:
    """Extract plain text from a supported document, or raise ``DocumentError``.

    Refuses an unsupported type, a macro-enabled Office format (by suffix or a
    ``vbaProject.bin`` member), an oversized file or a zip bomb; never executes
    anything. The returned text is untrusted and must be redacted before it
    reaches the model.
    """
    suffix = path.suffix.lower()
    if suffix in MACRO_SUFFIXES:
        msg = f"refusing a macro-enabled document type ({suffix}): {path.name}"
        raise DocumentError(msg)
    extractor = _EXTRACTORS.get(suffix)
    if extractor is None:
        msg = f"unsupported document type ({suffix or 'no suffix'}): {path.name}"
        raise DocumentError(msg)
    _check_size(path)
    return extractor(path)
